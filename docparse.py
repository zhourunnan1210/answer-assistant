# -*- coding: utf-8 -*-
"""简历/文档解析：把 PDF / Word / TXT / Markdown 提取为纯文本。

用于面试辅助：解析出的文本作为简历上下文随问题一起发给 LLM。
解析结果存进 config.json（resume_text），因此做长度收敛，避免配置膨胀。
"""
import os
import re

SUPPORTED_EXTS = (".pdf", ".docx", ".txt", ".md")
MAX_STORE_CHARS = 12000   # 存入配置的简历文本上限（约覆盖 3~5 页简历）


class DocParseError(Exception):
    pass


def _from_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(p for p in parts if p.strip())
    if not text.strip():
        raise DocParseError("PDF 未能提取到文字（可能是扫描件，请改用文字版 PDF 或 Word）")
    return text


def _from_docx(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # 简历常用表格排版，一并提取
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_text(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocParseError("文本文件编码无法识别（支持 UTF-8 / GBK）")


def extract_text(path: str) -> str:
    """提取文档纯文本：压缩空白、去 Markdown 标记、截断到存储上限。"""
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTS:
        raise DocParseError("不支持的格式：" + ext +
                            "（支持 PDF / Word / TXT / Markdown）")
    try:
        if ext == ".pdf":
            text = _from_pdf(path)
        elif ext == ".docx":
            text = _from_docx(path)
        else:
            text = _from_text(path)
    except DocParseError:
        raise
    except Exception as e:  # noqa: BLE001
        raise DocParseError(f"解析失败：{type(e).__name__}: {e}")

    text = text.replace("\r\n", "\n").replace("\r", "\n")  # 统一换行符
    text = re.sub(r"[ \t]+", " ", text)          # 压缩行内空白
    text = re.sub(r"\n{3,}", "\n\n", text)        # 压缩连续空行
    if ext == ".md":                              # 去掉 Markdown 排版标记
        text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
        text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.strip()
    if not text:
        raise DocParseError("文档内容为空")
    if len(text) > MAX_STORE_CHARS:
        text = text[:MAX_STORE_CHARS] + "\n……（内容过长，已截断）"
    return text
